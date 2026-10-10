from __future__ import annotations

import hashlib
import os
from pathlib import Path


MAX_FILE_BYTES = 20 * 1024 * 1024
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".py", ".ts", ".tsx", ".js", ".jsx", ".mjs", ".c", ".h",
    ".cc", ".cpp", ".hpp", ".cs", ".java", ".go", ".rs", ".sql", ".json", ".yaml",
    ".yml", ".toml", ".html", ".css", ".scss", ".sh", ".ps1", ".xml", ".ini", ".log",
}


def canonicalize_root(path: str) -> Path:
    if not isinstance(path, str) or not path.strip():
        raise ValueError("An existing folder path is required.")
    try:
        canonical = Path(path).expanduser().resolve(strict=True)
    except (OSError, RuntimeError) as error:
        raise ValueError("The selected folder does not exist or cannot be resolved.") from error
    if not canonical.is_dir():
        raise ValueError("The selected path must be a folder.")
    return canonical


def is_within_root(path: Path, root: Path) -> bool:
    try:
        canonical_root = root.resolve(strict=True)
        candidate = path.resolve(strict=False)
        common = Path(os.path.commonpath((str(canonical_root), str(candidate))))
        return os.path.normcase(str(common)) == os.path.normcase(str(canonical_root))
    except (OSError, RuntimeError, ValueError):
        return False


def extract_text(path: Path) -> str:
    try:
        canonical = path.resolve(strict=True)
    except (OSError, RuntimeError) as error:
        raise ValueError("The file does not exist or cannot be resolved.") from error
    if not canonical.is_file():
        raise ValueError("The selected path is not a file.")
    if canonical.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB indexing limit.")

    extension = canonical.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        return canonical.read_text(encoding="utf-8-sig", errors="replace")
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(canonical))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs cannot be indexed.")
        return "\n".join(text for page in reader.pages if (text := page.extract_text()))
    if extension == ".docx":
        from docx import Document

        document = Document(str(canonical))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {extension or '(no extension)'}.")


def chunk_text(text: str, *, max_chars: int = 1200, overlap: int = 160) -> list[str]:
    if not isinstance(max_chars, int) or isinstance(max_chars, bool) or max_chars <= 0:
        raise ValueError("max_chars must be a positive integer.")
    if not isinstance(overlap, int) or isinstance(overlap, bool) or overlap < 0 or overlap >= max_chars:
        raise ValueError("overlap must be a non-negative integer smaller than max_chars.")
    if not isinstance(text, str):
        raise ValueError("text must be a string.")
    if not text:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks


def content_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as content:
        for block in iter(lambda: content.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
